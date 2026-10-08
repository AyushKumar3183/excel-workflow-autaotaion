"""Workbook loader and validator for the workflow definition layer.

Reads workflow catalog sheets and step bindings from an Excel workbook, validates
the structural and relational integrity, and constructs validated WorkflowDefinition
objects.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple, Union

import openpyxl

from app.workflow.models import (
    Condition,
    DataSourceMetadata,
    InputSourceType,
    StepBinding,
    WorkflowDefinition,
    WorkflowInput,
    WorkflowStep,
)


class WorkflowLoadError(Exception):
    """Raised when the workbook file cannot be read or parsed."""


class WorkflowValidationError(Exception):
    """Raised when workflow data fails structural or relational validation."""


WORKFLOW_ID_PATTERN = re.compile(r"^WF[0-9A-Za-z_-]+$")


class WorkflowLoader:
    """Loads and validates workflows and step bindings from Excel workbooks."""

    REQUIRED_SHEETS = {"Workflows", "Step_Bindings"}
    
    WORKFLOW_COLUMNS = {
        "workflow_id",
        "workflow_name",
        "trigger",
        "inputs",
        "steps",
        "decision_logic",
        "tools_required",
        "expected_output",
    }

    BINDING_COLUMNS = {
        "workflow_id",
        "step_no",
        "tool",
        "params",
        "output_variable",
        "condition",
        "on_error",
    }

    def __init__(self, workbook_path: Optional[Union[str, Path]] = None):
        self.workbook_path = Path(workbook_path) if workbook_path else None

    def load(self, workbook_path: Optional[Union[str, Path]] = None) -> List[WorkflowDefinition]:
        """Load and validate all workflows from the given or configured workbook."""
        target_path = Path(workbook_path) if workbook_path else self.workbook_path
        if not target_path:
            raise WorkflowLoadError("No workbook path provided.")
        if not target_path.exists():
            raise WorkflowLoadError(f"Workbook file does not exist: {target_path}")

        sheets_data = self._read_workbook_sheets(target_path)
        return self.validate_and_build(sheets_data)

    def _read_workbook_sheets(self, path: Path) -> Dict[str, List[Dict[str, Any]]]:
        """Extract table rows as dictionaries per sheet."""
        # Try openpyxl first
        try:
            wb = openpyxl.load_workbook(filename=str(path), data_only=True)
            result: Dict[str, List[Dict[str, Any]]] = {}
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    result[sheetname] = []
                    continue
                
                # First non-empty row is header
                header_row_idx = None
                for idx, row in enumerate(rows):
                    if any(c is not None and str(c).strip() for c in row):
                        header_row_idx = idx
                        break
                
                if header_row_idx is None:
                    result[sheetname] = []
                    continue

                raw_headers = rows[header_row_idx]
                headers = [
                    str(h).strip().lower().replace(" ", "_") if h is not None else f"col_{i}"
                    for i, h in enumerate(raw_headers)
                ]

                sheet_records: List[Dict[str, Any]] = []
                for row in rows[header_row_idx + 1:]:
                    if not any(c is not None and str(c).strip() for c in row):
                        continue
                    record: Dict[str, Any] = {}
                    for col_idx, header in enumerate(headers):
                        val = row[col_idx] if col_idx < len(row) else None
                        record[header] = val
                    sheet_records.append(record)
                result[sheetname] = sheet_records
            return result

        except Exception as excel_err:
            # Check if file is a delimited text file (e.g. TSV / CSV)
            try:
                with open(path, "r", encoding="utf-8") as f:
                    sample = f.read(4096)
                    f.seek(0)
                    dialect = csv.Sniffer().sniff(sample, delimiters="\t,")
                    reader = csv.DictReader(f, dialect=dialect)
                    records = []
                    for r in reader:
                        normalized_r = {
                            k.strip().lower().replace(" ", "_"): v.strip() if isinstance(v, str) else v
                            for k, v in r.items()
                            if k is not None
                        }
                        records.append(normalized_r)
                # Text files only supply one sheet, usually 'Workflows'
                return {"Workflows": records}
            except Exception:
                raise WorkflowLoadError(
                    f"Could not open workbook {path} as Excel or delimited text: {excel_err}"
                ) from excel_err

    def validate_and_build(
        self, sheets_data: Dict[str, List[Dict[str, Any]]]
    ) -> List[WorkflowDefinition]:
        """Validate raw sheet data and construct WorkflowDefinition models."""
        # 1. Check for missing required sheets
        missing_sheets = self.REQUIRED_SHEETS - set(sheets_data.keys())
        if missing_sheets:
            raise WorkflowValidationError(
                f"Missing required sheet(s): {', '.join(sorted(missing_sheets))}"
            )

        workflow_rows = sheets_data.get("Workflows", [])
        binding_rows = sheets_data.get("Step_Bindings", [])

        if not workflow_rows:
            raise WorkflowValidationError("Sheet 'Workflows' contains no data rows.")

        # 2. Validate Workflow rows & collect workflow IDs
        known_workflows: Dict[str, Dict[str, Any]] = {}
        workflow_steps_text: Dict[str, List[str]] = {}

        for row_idx, row in enumerate(workflow_rows, start=2):
            raw_id = row.get("workflow_id")
            if not raw_id or not str(raw_id).strip():
                raise WorkflowValidationError(f"Row {row_idx} in 'Workflows' has missing or empty Workflow_ID.")
            
            wf_id = str(raw_id).strip()

            # Validate ID format
            if not WORKFLOW_ID_PATTERN.match(wf_id):
                raise WorkflowValidationError(
                    f"Invalid workflow ID format: '{wf_id}' at row {row_idx}. Must match pattern 'WF...'."
                )

            # Check for duplicate workflow IDs
            if wf_id in known_workflows:
                raise WorkflowValidationError(f"Duplicate workflow ID found: '{wf_id}' at row {row_idx}.")

            # Validate mandatory text fields
            wf_name = str(row.get("workflow_name") or "").strip()
            trigger = str(row.get("trigger") or "").strip()
            steps_desc = str(row.get("steps") or "").strip()

            if not wf_name:
                raise WorkflowValidationError(f"Workflow '{wf_id}' is missing required Workflow_Name.")
            if not trigger:
                raise WorkflowValidationError(f"Workflow '{wf_id}' is missing required Trigger.")
            if not steps_desc:
                raise WorkflowValidationError(f"Workflow '{wf_id}' is missing required Steps description.")

            # Parse steps description from business description (separated by →, ->, or ;)
            parsed_step_names = [
                s.strip()
                for s in re.split(r"→|->|;", steps_desc)
                if s.strip()
            ]
            if not parsed_step_names:
                raise WorkflowValidationError(f"Workflow '{wf_id}' has no valid steps described.")

            workflow_steps_text[wf_id] = parsed_step_names
            known_workflows[wf_id] = row

        # 3. Validate Step_Bindings rows
        bindings_by_wf: Dict[str, List[StepBinding]] = {wf_id: [] for wf_id in known_workflows}

        for row_idx, row in enumerate(binding_rows, start=2):
            raw_id = row.get("workflow_id")
            if not raw_id or not str(raw_id).strip():
                raise WorkflowValidationError(f"Row {row_idx} in 'Step_Bindings' is missing workflow_id.")
            
            wf_id = str(raw_id).strip()

            # Check for bindings referencing unknown workflows
            if wf_id not in known_workflows:
                raise WorkflowValidationError(
                    f"Step binding at row {row_idx} references unknown workflow: '{wf_id}'."
                )

            # Validate step number
            raw_step_no = row.get("step_no")
            if raw_step_no is None:
                raise WorkflowValidationError(
                    f"Missing step number at row {row_idx} for workflow '{wf_id}'."
                )
            try:
                step_no = int(raw_step_no)
            except (ValueError, TypeError):
                raise WorkflowValidationError(
                    f"Invalid step number '{raw_step_no}' at row {row_idx} for workflow '{wf_id}'."
                )

            if step_no < 1:
                raise WorkflowValidationError(
                    f"Invalid step number {step_no} at row {row_idx} for workflow '{wf_id}'. Step numbers must be >= 1."
                )

            # Check if step number exceeds defined steps for workflow
            defined_step_count = len(workflow_steps_text[wf_id])
            if step_no > defined_step_count:
                raise WorkflowValidationError(
                    f"Step binding references unknown step {step_no} for workflow '{wf_id}' "
                    f"(workflow defines {defined_step_count} steps)."
                )

            # Validate tool
            tool = str(row.get("tool") or "").strip()
            if not tool:
                raise WorkflowValidationError(
                    f"Missing required tool name in step binding at row {row_idx} for workflow '{wf_id}'."
                )

            # Validate output_variable
            output_var = str(row.get("output_variable") or "").strip()
            if not output_var:
                raise WorkflowValidationError(
                    f"Missing required output_variable in step binding at row {row_idx} for workflow '{wf_id}'."
                )

            # Parse params (can be dict or JSON string)
            raw_params = row.get("params")
            params_dict: Dict[str, Any] = {}
            if raw_params:
                if isinstance(raw_params, dict):
                    params_dict = raw_params
                elif isinstance(raw_params, str) and raw_params.strip():
                    try:
                        params_dict = json.loads(raw_params.strip())
                        if not isinstance(params_dict, dict):
                            raise ValueError("Params JSON must be an object.")
                    except Exception as e:
                        raise WorkflowValidationError(
                            f"Invalid params JSON in step binding row {row_idx} for '{wf_id}': {e}"
                        )
                else:
                    raise WorkflowValidationError(
                        f"Unsupported params type in step binding row {row_idx} for '{wf_id}': {type(raw_params)}"
                    )

            # Parse condition
            raw_condition = row.get("condition")
            condition_obj: Optional[Condition] = None
            if raw_condition:
                try:
                    condition_obj = Condition.from_raw(raw_condition)
                except Exception as e:
                    raise WorkflowValidationError(
                        f"Invalid condition in step binding row {row_idx} for '{wf_id}': {e}"
                    )

            # Parse on_error
            raw_on_error = str(row.get("on_error") or "FAIL").strip().upper()
            try:
                binding = StepBinding(
                    workflow_id=wf_id,
                    step_no=step_no,
                    tool=tool,
                    params=params_dict,
                    output_variable=output_var,
                    condition=condition_obj,
                    on_error=raw_on_error,
                )
            except Exception as e:
                raise WorkflowValidationError(
                    f"Validation error in step binding row {row_idx} for '{wf_id}': {e}"
                )

            bindings_by_wf[wf_id].append(binding)

        # 4. Parse optional Data_Sources sheet
        data_sources_rows = sheets_data.get("Data_Sources", [])
        data_sources_by_wf: Dict[str, Dict[str, DataSourceMetadata]] = {wf_id: {} for wf_id in known_workflows}

        for row in data_sources_rows:
            wf_id = str(row.get("workflow_id") or "").strip()
            if not wf_id or wf_id not in known_workflows:
                continue
            input_name = str(row.get("input_name") or "").strip()
            if not input_name:
                continue
            raw_source_type = str(row.get("source_type") or "user").strip().lower()
            try:
                source_type = InputSourceType(raw_source_type)
            except ValueError:
                source_type = InputSourceType.USER

            raw_ds = row.get("data_source")
            data_source = str(raw_ds).strip() if raw_ds is not None and str(raw_ds).strip() else None

            default_val = row.get("default_value")
            raw_req = row.get("required")
            if isinstance(raw_req, bool):
                is_req = raw_req
            elif isinstance(raw_req, str):
                is_req = raw_req.strip().lower() in {"true", "1", "yes"}
            else:
                is_req = True if raw_req is None else bool(raw_req)

            desc = str(row.get("description") or "").strip() or None

            ds_meta = DataSourceMetadata(
                name=input_name,
                source_type=source_type,
                data_source=data_source,
                default_value=default_val,
                required=is_req,
                description=desc,
            )
            data_sources_by_wf[wf_id][input_name] = ds_meta

        # 5. Validate completeness of step bindings for each workflow
        workflow_definitions: List[WorkflowDefinition] = []

        for wf_id, wf_row in known_workflows.items():
            wf_bindings = bindings_by_wf[wf_id]

            # Check for missing required step bindings
            if not wf_bindings:
                raise WorkflowValidationError(
                    f"Missing required step bindings for workflow '{wf_id}'."
                )

            # Check step sequence completeness and duplicate step numbers
            sorted_bindings = sorted(wf_bindings, key=lambda b: b.step_no)
            seen_steps: Set[int] = set()
            for b in sorted_bindings:
                if b.step_no in seen_steps:
                    raise WorkflowValidationError(
                        f"Duplicate step number {b.step_no} in bindings for workflow '{wf_id}'."
                    )
                seen_steps.add(b.step_no)

            expected_count = len(workflow_steps_text[wf_id])
            if len(sorted_bindings) != expected_count:
                missing = set(range(1, expected_count + 1)) - seen_steps
                raise WorkflowValidationError(
                    f"Missing step bindings for workflow '{wf_id}': step(s) {sorted(missing)} not bound. "
                    f"Expected {expected_count} steps, got {len(sorted_bindings)}."
                )

            # Construct WorkflowStep objects
            workflow_steps: List[WorkflowStep] = []
            for b in sorted_bindings:
                step_title = workflow_steps_text[wf_id][b.step_no - 1]
                workflow_steps.append(
                    WorkflowStep(
                        step_no=b.step_no,
                        step_name=step_title,
                        description=f"{step_title} using {b.tool}",
                        tool=b.tool,
                        params=b.params,
                        output_variable=b.output_variable,
                        condition=b.condition,
                        on_error=b.on_error,
                    )
                )

            # Parse inputs list from Inputs column
            raw_inputs = str(wf_row.get("inputs") or "")
            parsed_inputs = self._parse_inputs(raw_inputs)

            # Parse tools_required list
            raw_tools = str(wf_row.get("tools_required") or "")
            tools_list = [t.strip() for t in re.split(r";|,", raw_tools) if t.strip()]

            wf_ds = data_sources_by_wf.get(wf_id, {})

            # Build WorkflowDefinition
            try:
                definition = WorkflowDefinition(
                    workflow_id=wf_id,
                    name=str(wf_row.get("workflow_name")).strip(),
                    trigger=str(wf_row.get("trigger")).strip(),
                    inputs=parsed_inputs,
                    steps=workflow_steps,
                    data_sources=wf_ds,
                    decision_logic=str(wf_row.get("decision_logic") or "").strip() or None,
                    tools_required=tools_list,
                    expected_output=str(wf_row.get("expected_output") or "").strip() or None,
                    metadata={"source_row": wf_row},
                )
                workflow_definitions.append(definition)
            except Exception as e:
                raise WorkflowValidationError(
                    f"Error creating WorkflowDefinition for '{wf_id}': {e}"
                )

        return workflow_definitions

    @staticmethod
    def _parse_inputs(raw_inputs: str) -> List[WorkflowInput]:
        """Convert input string (e.g. 'Product inventory CSV; minimum stock threshold') into WorkflowInput objects."""
        if not raw_inputs or not raw_inputs.strip():
            return []
        
        parts = [p.strip() for p in re.split(r";|,", raw_inputs) if p.strip()]
        inputs: List[WorkflowInput] = []

        for p in parts:
            name_slug = re.sub(r"[^a-zA-Z0-9_]+", "_", p.lower()).strip("_")
            input_type = "string"
            if any(kw in p.lower() for kw in ("csv", "xlsx", "file", "spreadsheet")):
                input_type = "file"
            elif any(kw in p.lower() for kw in ("threshold", "rate", "percent", "quantity", "price")):
                input_type = "float"
            elif any(kw in p.lower() for kw in ("list", "catalog", "logs")):
                input_type = "list"
            elif any(kw in p.lower() for kw in ("id", "email", "name")):
                input_type = "string"

            inputs.append(
                WorkflowInput(
                    name=name_slug or "input_param",
                    description=p,
                    type=input_type,
                    required=True,
                )
            )

        return inputs


def load_workbook(workbook_path: Union[str, Path]) -> List[WorkflowDefinition]:
    """Convenience helper to load and validate workflows from an Excel workbook."""
    loader = WorkflowLoader(workbook_path)
    return loader.load()
